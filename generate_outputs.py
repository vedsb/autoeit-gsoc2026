"""
generate_outputs.py — Pre-generates all AutoEIT outputs for the notebook.
Run: python generate_outputs.py
"""
import os, re, unicodedata
from collections import Counter
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
import pandas as pd
import matplotlib
matplotlib.use("Agg")  # non-interactive backend
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
import spacy

print("Loading spaCy model...")
NLP = spacy.load("es_core_news_sm")

INPUT_XLSX = os.path.join(
    "AutoEIT Test Files", "Sample Audio Files and Transcriptions",
    "AutoEIT Sample Transcriptions for Scoring.xlsx"
)
OUTPUT_XLSX = "AutoEIT_Scored_Output.xlsx"

# ── Copied helpers from eit_scorer.py ─────────────────────────────────────────
GARBLE_RE = re.compile(
    r"\b(xxx|xx|x\b|gibberish|pause|no response|inaudible|unintelligible)\b"
    r"|\[.*?\]|\.\.\.", re.IGNORECASE,
)
FUNCTION_WORDS = {
    "a","al","ante","bajo","con","contra","de","del","desde","durante","en",
    "entre","hacia","hasta","mediante","para","por","segun","sin","sobre","tras",
    "el","la","los","las","un","una","unos","unas","y","e","o","u","pero","sino",
    "ni","mas","que","se","me","te","le","lo","les","nos","os","quien","quienes",
    "cual","cuales","donde","cuando","como","no","si","ya","aun","tambien",
    "tampoco","yo","tu","el","ella","nosotros","vosotros","ellos","ellas",
    "usted","ustedes","este","esta","estos","estas","ese","esa","esos","esas",
    "aquel","aquella","aquellos","aquellas","mi","tu","su","sus","mis","tus",
    "mio","mia","tuyo","tuya","suyo","suya","es","son","esta","estan",
}
SYNONYMS = {"pero": "y", "e": "y", "u": "o", "muy": ""}

def normalize(text):
    text = text.lower()
    nfkd = unicodedata.normalize("NFKD", text)
    text = "".join(c for c in nfkd if not unicodedata.combining(c))
    text = re.sub(r"[^\w\s]", " ", text)
    return re.sub(r"\s+", " ", text).strip()

def strip_word_count(s):
    return re.sub(r"\s*\(\d+\)\s*$", "", s).strip()

def extract_best_final_response(t):
    if not t or not t.strip(): return ""
    t = re.sub(r"\[.*?\]", " ", t)
    t = re.sub(r"\b\w*-\s+", " ", t)
    if "- " in t:
        parts = re.split(r"\w+-\s+", t)
        t = parts[-1] if len(parts) > 1 else t
    t = GARBLE_RE.sub(" ", t)
    return re.sub(r"\s+", " ", t).strip()

def get_content_words(text):
    doc = NLP(text)
    return [
        normalize(tok.lemma_) for tok in doc
        if not tok.is_space and not tok.is_punct
        and tok.pos_ in {"NOUN","VERB","ADJ","ADV","PROPN"}
        and normalize(tok.lemma_) not in FUNCTION_WORDS
    ]

def is_all_garbled(text):
    return not any(len(w) > 1 for w in text.split())

def compute_overlap(stim_c, trans_c):
    if not stim_c: return 1.0, 0, 0
    sc = Counter(stim_c)
    matched = sum(min(cnt, trans_c.count(w)) for w, cnt in sc.items())
    return matched/len(stim_c), matched, len(stim_c)

def word_order_preserved(stim_c, trans_c):
    idx = 0
    for w in stim_c:
        while idx < len(trans_c) and trans_c[idx] != w: idx += 1
        if idx >= len(trans_c): return False
        idx += 1
    return True

def apply_synonyms(words):
    return [SYNONYMS.get(w,w) for w in words if SYNONYMS.get(w,w)]

def score_sentence(stim_raw, trans_raw):
    stim_clean = strip_word_count(stim_raw or "")
    trans_clean = extract_best_final_response(trans_raw or "")
    stim_norm = normalize(stim_clean)
    trans_norm = normalize(trans_clean)
    details = {}
    if is_all_garbled(trans_norm):
        return 0, {**details, "decision":"Score 0: empty/unintelligible"}
    trans_words = trans_norm.split()
    if stim_norm == trans_norm:
        return 4, {**details, "decision":"Score 4: exact match"}
    stim_mapped  = apply_synonyms(stim_norm.split())
    trans_mapped = apply_synonyms(trans_norm.split())
    if stim_mapped == trans_mapped:
        return 4, {**details, "decision":"Score 4: synonym-only substitution"}
    stim_c  = get_content_words(stim_clean)
    trans_c = get_content_words(trans_clean) if trans_clean else []
    ratio, matched, total = compute_overlap(stim_c, trans_c)
    order_ok = word_order_preserved(stim_c, trans_c)
    effective = ratio + (0.10 if order_ok else 0.0)
    details = dict(idea_unit_ratio=round(ratio,3), effective_ratio=round(effective,3),
                   matched=matched, total=total, word_order=order_ok,
                   stim_content=stim_c, trans_content=trans_c)
    only_function = all(w in FUNCTION_WORDS for w in trans_words)
    if len(trans_words) <= 1 or only_function:
        return 1, {**details, "decision":"Score 1: ≤1 word or only function words"}
    if matched <= 1 and total > 2:
        return 1, {**details, "decision":"Score 1: ≤1 content word matched"}
    stim_m2  = apply_synonyms([normalize(t.lemma_) for t in NLP(stim_clean) if not t.is_space and not t.is_punct])
    trans_m2 = apply_synonyms([normalize(t.lemma_) for t in NLP(trans_clean) if not t.is_space and not t.is_punct and trans_clean])
    if stim_m2 == trans_m2:
        return 3, {**details, "decision":"Score 3: synonym-adjusted lemma match"}
    if effective >= 0.75:
        return 3, {**details, "decision":f"Score 3: effective ratio {effective:.2f} ≥ 0.75"}
    elif effective >= 0.40:
        return 2, {**details, "decision":f"Score 2: effective ratio {effective:.2f}"}
    else:
        return 1, {**details, "decision":f"Score 1: effective ratio {effective:.2f} < 0.40"}

# ── Score all ─────────────────────────────────────────────────────────────────
COLOURS = {0:"FFD7D7",1:"FFE8C8",2:"FFFACD",3:"D5F5D5",4:"C8E6FA"}
wb = openpyxl.load_workbook(INPUT_XLSX)
all_rows = []

for sname in wb.sheetnames:
    if sname == "Info": continue
    ws = wb[sname]
    for row in ws.iter_rows(min_row=2):
        if row[0].value is None or not isinstance(row[0].value,(int,float)): continue
        stim  = row[1].value or ""
        trans = row[2].value or ""
        sc, det = score_sentence(stim, trans)
        row[3].value = sc
        fill = COLOURS.get(sc,"FFFFFF")
        row[3].fill = PatternFill(start_color=fill,end_color=fill,fill_type="solid")
        row[3].font = Font(bold=True)
        row[3].alignment = Alignment(horizontal="center")
        all_rows.append(dict(
            participant=sname, sentence=int(row[0].value),
            stimulus=strip_word_count(stim), transcription=trans,
            score=sc, decision=det.get("decision",""),
            idea_unit_ratio=det.get("idea_unit_ratio",None),
        ))

wb.save(OUTPUT_XLSX)
print(f"✅ Scored output: {OUTPUT_XLSX}")

df = pd.DataFrame(all_rows)

# ── Statistics ────────────────────────────────────────────────────────────────
print("\n=== Per-Participant Summary ===")
summary = df.groupby("participant")["score"].agg(["count","sum","mean","std"])
summary.columns = ["N","Sum","Mean","SD"]
summary["% of Max"] = (summary["Sum"]/(summary["N"]*4)*100).round(1)
print(summary.round(2).to_string())
print(f"\nOverall: Mean={df.score.mean():.3f}, SD={df.score.std():.3f}")

# ── Chart ─────────────────────────────────────────────────────────────────────
COLOUR_MAP = {0:"#FFAAAA",1:"#FFC878",2:"#FFF066",3:"#78E878",4:"#78C8F0"}
participants = df["participant"].unique().tolist()

fig, axes = plt.subplots(1, len(participants)+1, figsize=(5*(len(participants)+1), 4))

for i, pname in enumerate(participants):
    ax = axes[i]
    pscores = df[df.participant==pname]["score"].tolist()
    dist = Counter(pscores)
    bars = ax.bar([str(s) for s in range(5)],[dist.get(s,0) for s in range(5)],
                  color=[COLOUR_MAP[s] for s in range(5)],edgecolor="grey",linewidth=0.7)
    ax.set_title(f"Participant\n{pname}",fontsize=10,fontweight="bold")
    ax.set_xlabel("Score",fontsize=9)
    ax.set_ylabel("Count" if i==0 else "",fontsize=9)
    ax.set_ylim(0,30)
    for bar,v in zip(bars,[dist.get(s,0) for s in range(5)]):
        if v: ax.text(bar.get_x()+bar.get_width()/2,bar.get_height()+0.3,str(v),
                      ha="center",va="bottom",fontsize=9,fontweight="bold")
    ax.text(0.97,0.97,f"Mean={sum(pscores)/len(pscores):.2f}",transform=ax.transAxes,
            ha="right",va="top",fontsize=8,color="dimgrey")

ax = axes[-1]
all_s = df["score"].tolist()
dist_all = Counter(all_s)
bars = ax.bar([str(s) for s in range(5)],[dist_all.get(s,0) for s in range(5)],
              color=[COLOUR_MAP[s] for s in range(5)],edgecolor="grey",linewidth=0.7)
ax.set_title("Overall\n(All Participants)",fontsize=10,fontweight="bold")
ax.set_xlabel("Score",fontsize=9)
ax.set_ylim(0,65)
for bar,v in zip(bars,[dist_all.get(s,0) for s in range(5)]):
    if v: ax.text(bar.get_x()+bar.get_width()/2,bar.get_height()+0.5,str(v),
                  ha="center",va="bottom",fontsize=9,fontweight="bold")

legend_patches = [mpatches.Patch(color=COLOUR_MAP[s],label=f"Score {s}") for s in range(5)]
fig.legend(handles=legend_patches,loc="lower center",ncol=5,fontsize=9,
           framealpha=0.8,title="Score",title_fontsize=9)
plt.suptitle("EIT Score Distributions by Participant\n(Ortega 2000 Meaning-Based Rubric)",
             fontsize=12,fontweight="bold",y=1.02)
plt.tight_layout()
plt.savefig("score_distributions.png",dpi=150,bbox_inches="tight")
print("✅ Chart saved: score_distributions.png")

# ── Sentence-level detail table ───────────────────────────────────────────────
print("\n=== Sentence-Level Scores (all participants) ===")
for pname in participants:
    pdf = df[df.participant==pname].copy()
    print(f"\n{pname}  (Sum={pdf.score.sum()}/120, Mean={pdf.score.mean():.2f})")
    print(f"  {'#':>2}  {'Sc':>2}  {'Ratio':>5}  Stimulus")
    print(f"  {'─'*2}  {'─'*2}  {'─'*5}  {'─'*48}")
    for _, r in pdf.iterrows():
        stim = r["stimulus"][:46]
        ratio = f"{r['idea_unit_ratio']:.2f}" if r["idea_unit_ratio"] is not None else "  —  "
        print(f"  {r['sentence']:>2}  {r['score']:>2}  {ratio:>5}  {stim}")

# ── Validation cases ──────────────────────────────────────────────────────────
print("\n=== Manual Validation Against Rubric Examples ===")
VALIDATION_CASES = [
    ("Quiero cortarme el pelo",     "Quiero cortarme el pelo",            4, "Exact"),
    ("Quiero cortarme el pelo",     "Quiero cortar mi pelo",              3, "Meaning preserved"),
    ("El carro lo tiene Pedro",     "el carro tiene Pedro",               2, "Missing clitic 'lo'"),
    ("Dudo que sepa manejar muy bien","dudo/tu no? sepiar exx muy bien",  1, "Mostly garbled"),
    ("Ella sólo bebe cerveza y no come nada","Ella sola cerveza y no come nada",2,"Partial meaning"),
    ("Después de cenar me fui a dormir tranquilo","Después de cenar me fui a dormir tranquilo",4,"Exact"),
    ("Quiero una casa en la que vivan mis animales","Quiero una casa en que viven mis animales",3,"Minor grammar"),
    ("El niño al que se le murió el gato está triste","El niño se murió el gato es triste",3,"Full meaning, grammar errors"),
]
print(f"  {'#':>2}  {'Exp':>4}  {'Got':>4}  Match  Stimulus")
print(f"  {'─'*2}  {'─'*4}  {'─'*4}  {'─'*5}  {'─'*45}")
all_match = True
for i,(stim,trans,exp,note) in enumerate(VALIDATION_CASES,1):
    sc,_ = score_sentence(stim,trans)
    ok = "✅" if sc==exp else "❌"
    if sc!=exp: all_match=False
    print(f"  {i:>2}  {exp:>4}  {sc:>4}  {ok}     {stim[:43]}")
print(f"\n  {'✅ All validation cases pass!' if all_match else '⚠️  Some cases differ — see above'}")

print("\n✅ All outputs generated successfully.")
print("   - AutoEIT_Scored_Output.xlsx")
print("   - score_distributions.png")
